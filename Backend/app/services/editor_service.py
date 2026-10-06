"""
Editor Service: Parse DOCX documents into layout structures for TemplateDocumentEditor and apply modifications.
"""

from pathlib import Path
from typing import Dict, Any, List, Optional
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

from app.core.config import settings
from app.core.logging import logger
from app.infrastructure.storage.local_storage import LocalStorageProvider


class EditorService:
    def __init__(self, storage_provider: Optional[LocalStorageProvider] = None):
        self.storage = storage_provider or LocalStorageProvider()

    def _parse_doc_to_layout(self, doc: docx.Document, filename: str) -> Dict[str, Any]:
        """Core parser: Iterates over doc.element.body in sequential order for authentic Word document layout."""
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        page = {
            "page_width": 595.3,
            "page_height": 841.9,
            "top_margin": 54.0,
            "right_margin": 54.0,
            "bottom_margin": 54.0,
            "left_margin": 54.0,
        }

        if doc.sections:
            sec = doc.sections[0]
            if sec.page_width:
                page["page_width"] = round(sec.page_width.pt, 1)
            if sec.page_height:
                page["page_height"] = round(sec.page_height.pt, 1)
            if sec.top_margin:
                page["top_margin"] = round(sec.top_margin.pt, 1)
            if sec.right_margin:
                page["right_margin"] = round(sec.right_margin.pt, 1)
            if sec.bottom_margin:
                page["bottom_margin"] = round(sec.bottom_margin.pt, 1)
            if sec.left_margin:
                page["left_margin"] = round(sec.left_margin.pt, 1)

        def get_text_align(alignment, text: str = "") -> str:
            if alignment == WD_ALIGN_PARAGRAPH.CENTER:
                return "center"
            elif alignment == WD_ALIGN_PARAGRAPH.RIGHT:
                return "right"
            elif alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                return "justify"

            # Contextual fallback for Tamil administrative formats
            txt = text.strip()
            if any(txt.startswith(h) for h in [
                "ஈரோடு மாவட்ட ஆட்சித் தலைவர்",
                "மாவட்ட நிர்வாக நடுவர்",
                "முன்னிலை:",
                "// அலுவலகக் குறிப்பு //",
                "// குறிப்பாணை //",
                "-------"
            ]):
                return "center"
            elif any(txt.startswith(s) for s in [
                "மாவட்ட ஆட்சித் தலைவர்,",
                "மாவட்ட ஆட்சித் தலைவருக்காக",
                "மாவட்ட ஆட்சியரின் நேர்முக உதவியாளர்",
                "ஈரோடு."
            ]) and not txt.startswith("ஈரோடு மாவட்ட ஆட்சித் தலைவர்"):
                return "right"
            elif len(txt) > 70 and not (
                txt.startswith("பெறுநர்:") or txt.startswith("நகல்:") or txt.startswith("இணைப்பு:") or txt.startswith("ந.க.")
            ):
                return "justify"

            return "left"

        def parse_paragraph(p, pid: str, is_table_cell: bool = False) -> Dict[str, Any]:
            txt = p.text or ""
            align = get_text_align(p.alignment, txt)
            runs_data = []
            for r in p.runs:
                if not r.text:
                    continue
                r_style: Dict[str, Any] = {}
                if r.bold:
                    r_style["fontWeight"] = "bold"
                if r.italic:
                    r_style["fontStyle"] = "italic"
                if r.underline:
                    r_style["textDecoration"] = "underline"
                if r.font and r.font.size:
                    r_style["fontSize"] = f"{r.font.size.pt}pt"
                runs_data.append({
                    "text": r.text,
                    "style": r_style
                })

            if not runs_data:
                runs_data = [{"text": txt, "style": {}}]

            # Indent only for narrative body paragraphs (not headers, labels, table cells, copies, or signatory)
            is_body_para = (align == "justify" and not is_table_cell and not any(
                txt.strip().startswith(prefix) for prefix in [
                    "பெறுநர்:", "நகல்:", "இணைப்பு:", "ந.க.", "பொருள்:", "பார்வை:", "உத்தரவு:", "பணிந்தனுப்பப்படுகிறது:"
                ]
            ))
            text_indent = "2rem" if (is_body_para or (p.paragraph_format.first_line_indent and not is_table_cell and align not in ("center", "right"))) else "0"

            p_style: Dict[str, Any] = {
                "textAlign": align,
                "fontFamily": "'TAU-Marutham', 'Noto Sans Tamil', 'Plus Jakarta Sans', sans-serif",
                "lineHeight": "1.6",
                "margin": "0 0 0.5rem 0",
                "color": "#102C57",
                "textIndent": text_indent
            }

            return {
                "id": pid,
                "type": "paragraph",
                "text": txt,
                "runs": runs_data,
                "style": p_style
            }

        blocks: List[Dict[str, Any]] = []
        p_count = 0
        t_count = 0

        # Sequential iteration over document body elements in exact order
        for elem in doc.element.body:
            if elem.tag.endswith("p"):
                p = Paragraph(elem, doc)
                p_count += 1
                blocks.append(parse_paragraph(p, f"p_{p_count}"))
            elif elem.tag.endswith("tbl"):
                t_count += 1
                table = Table(elem, doc)
                table_rows = []
                for r_idx, row in enumerate(table.rows):
                    row_cells = []
                    for c_idx, cell in enumerate(row.cells):
                        cell_blocks = []
                        for cp_idx, cp in enumerate(cell.paragraphs):
                            p_count += 1
                            cell_blocks.append(parse_paragraph(cp, f"p_t{t_count}_r{r_idx}_c{c_idx}_{cp_idx+1}", is_table_cell=True))
                        row_cells.append({
                            "colSpan": 1,
                            "width": f"{round(100.0 / max(1, len(row.cells)), 1)}%",
                            "blocks": cell_blocks
                        })
                    table_rows.append(row_cells)
                blocks.append({
                    "id": f"table_{t_count}",
                    "type": "table",
                    "rows": table_rows
                })

        return {
            "filename": filename,
            "revision": 1,
            "page": page,
            "blocks": blocks
        }

    def docx_to_layout(self, file_path: Path, filename: str) -> Dict[str, Any]:
        """Parses a Word (.docx) document into a block-based JSON layout for interactive editing."""
        doc = docx.Document(str(file_path))
        return self._parse_doc_to_layout(doc, filename)

    def bytes_to_layout(self, raw_bytes: bytes, filename: str) -> Dict[str, Any]:
        """Parses in-memory DOCX binary bytes into an interactive block layout."""
        import io
        stream = io.BytesIO(raw_bytes)
        doc = docx.Document(stream)
        return self._parse_doc_to_layout(doc, filename)

    def template_model_to_layout(self, tpl) -> Dict[str, Any]:
        """Converts a DocumentTemplate database model into an interactive block layout with proper alignments and keyword placeholders."""
        # 1. If custom block layout data was previously saved by the user in DB, return it directly
        if getattr(tpl, "template_data", None) and isinstance(tpl.template_data, dict) and "blocks" in tpl.template_data:
            return tpl.template_data

        blocks = []
        code = getattr(tpl, "template_code", "") or ""
        category = getattr(tpl, "category", "") or ""

        if "note" in code or category == "INTERNAL_NOTE":
            # =========================================================================
            # OFFICE NOTE LAYOUT
            # =========================================================================
            blocks.append({
                "id": "p_heading",
                "type": "paragraph",
                "text": getattr(tpl, "heading_prefix", None) or "// அலுவலகக் குறிப்பு //",
                "runs": [{"text": getattr(tpl, "heading_prefix", None) or "// அலுவலகக் குறிப்பு //", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "center", "fontSize": "12pt", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.8rem 0"}
            })

            subj = getattr(tpl, "subject_template", None) or "வருவாய் வசூல் சட்டம் – {{statute_cited}} – {{district_name}} மாவட்டம், {{taluk_name}} வட்டம், {{street_and_locality}} என்ற முகவரியில் {{living_verb}} {{defaulter_name}} என்ற {{entity_label}} அரசிற்கு செலுத்த வேண்டிய {{dues_label}} மொத்தம் {{total_amount}}ஐ வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்ய கோரியது - உத்திரவிடுதல் - தொடர்பாக."
            ref = getattr(tpl, "reference_template", None) or "{{reference_text}}"

            blocks.append({
                "id": "table_subj_ref",
                "type": "table",
                "rows": [
                    [
                        {"colSpan": 1, "width": "15%", "blocks": [{"id": "p_subj_lbl", "type": "paragraph", "text": "பொருள்:", "runs": [{"text": "பொருள்:", "style": {"fontWeight": "bold"}}], "style": {"textAlign": "left", "fontSize": "11pt"}}]},
                        {"colSpan": 1, "width": "85%", "blocks": [{"id": "p_subj_val", "type": "paragraph", "text": subj, "runs": [{"text": subj}], "style": {"textAlign": "left", "fontSize": "11pt"}}]}
                    ],
                    [
                        {"colSpan": 1, "width": "15%", "blocks": [{"id": "p_ref_lbl", "type": "paragraph", "text": "பார்வை:", "runs": [{"text": "பார்வை:", "style": {"fontWeight": "bold"}}], "style": {"textAlign": "left", "fontSize": "11pt"}}]},
                        {"colSpan": 1, "width": "85%", "blocks": [{"id": "p_ref_val", "type": "paragraph", "text": ref, "runs": [{"text": ref}], "style": {"textAlign": "left", "fontSize": "11pt"}}]}
                    ]
                ]
            })

            blocks.append({
                "id": "p_note_lbl",
                "type": "paragraph",
                "text": "பணிந்தனுப்பப்படுகிறது:",
                "runs": [{"text": "பணிந்தனுப்பப்படுகிறது:", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "left", "fontSize": "11pt", "margin": "0.6rem 0 0.4rem 0"}
            })

            p1 = getattr(tpl, "order_para1_template", None) or "{{district_name}} மாவட்டம், {{taluk_name}} வட்டம், {{street_and_locality}} என்ற முகவரியில் {{living_verb}} {{defaulter_name}} என்ற {{entity_label}} அரசிற்கு செலுத்த வேண்டிய {{dues_label}} மொத்தம் {{total_amount}}ஐ வருவாய் வசூல் சட்டத்தின் கீழ் வசூலிக்குமாறு பார்வையில் காணும் கடிதத்தில் தெரிவிக்கப்பட்டுள்ளது."
            blocks.append({
                "id": "p_order_para1",
                "type": "paragraph",
                "text": p1,
                "runs": [{"text": p1}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "textIndent": "2rem", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.6rem 0", "lineHeight": "1.6"}
            })

            p2 = getattr(tpl, "order_para2_template", None) or "மேற்படி முகவரியில் {{living_verb}} {{defaulter_name}} என்பவரின் அசையும் மற்றும் அசையா சொத்துகளிலிருந்து அரசிற்கு செலுத்த வேண்டிய {{dues_label}} மொத்தம் {{total_amount}}ஐ ({{amount_in_tamil_words}}) வருவாய் வசூல் சட்டப்படி வசூல் செய்து “{{dd_favour_of}}” என்ற பெயரில் வங்கி வரைவோலையாக (Demand Draft) எடுத்து {{dispatch_address}} என்ற அலுவலகத்திற்கு அசல் வங்கி வரைவோலையினை அனுப்பி அதன் விவரத்தினை நகல் வங்கி வரைவோலையுடன் இவ்வலுவலகத்திற்கு அனுப்பி வைக்குமாறு {{taluk_name}} வருவாய் வட்டாட்சியருக்கு தெரிவிக்கலாம்."
            blocks.append({
                "id": "p_order_para2",
                "type": "paragraph",
                "text": p2,
                "runs": [{"text": p2}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "textIndent": "2rem", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.6rem 0", "lineHeight": "1.6"}
            })

            p3 = getattr(tpl, "order_para3_template", None) or "எனவே மேற்படி தொகையை வருவாய் நிலை ஆணை எண் 41 மற்றும் வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூல் செய்ய {{taluk_name}} வருவாய் வட்டாட்சியருக்கு அதிகாரம் வழங்கி இதன் மூலம் உத்தரவிடலாம்."
            blocks.append({
                "id": "p_order_para3",
                "type": "paragraph",
                "text": p3,
                "runs": [{"text": p3}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "textIndent": "2rem", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.6rem 0", "lineHeight": "1.6"}
            })

            submission = getattr(tpl, "signatory_text", None) or "உத்தரவினை எதிர்நோக்கி செயல்முறை வரைவு ஒப்புதலுக்காக மாவட்ட ஆட்சித்தலைவர் அவர்களுக்கு பணிவுடன் சமர்ப்பிக்கப்படுகிறது."
            blocks.append({
                "id": "p_submission",
                "type": "paragraph",
                "text": submission,
                "runs": [{"text": submission, "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "fontFamily": "'TAU-Marutham', serif", "margin": "1rem 0"}
            })

        else:
            # =========================================================================
            # PROCEEDINGS / GENERAL PROCEEDINGS LAYOUT
            # =========================================================================
            blocks.append({
                "id": "p_heading",
                "type": "paragraph",
                "text": "ஈரோடு மாவட்ட ஆட்சித் தலைவர் மற்றும்\nமாவட்ட நிர்வாக நடுவர் அவர்களின் செயல்முறைகள், ஈரோடு",
                "runs": [{"text": "ஈரோடு மாவட்ட ஆட்சித் தலைவர் மற்றும்\nமாவட்ட நிர்வாக நடுவர் அவர்களின் செயல்முறைகள், ஈரோடு", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "center", "fontSize": "12pt", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.4rem 0"}
            })

            blocks.append({
                "id": "p_collector",
                "type": "paragraph",
                "text": "முன்னிலை: {{collector_name}}",
                "runs": [{"text": "முன்னிலை: {{collector_name}}", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "center", "fontSize": "11pt", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.5rem 0"}
            })

            blocks.append({
                "id": "p_roc",
                "type": "paragraph",
                "text": "ந.க. {{nk_no}}/{{doc_year}}/{{section_code}}                                            நாள்: {{doc_date}}",
                "runs": [{"text": "ந.க. {{nk_no}}/{{doc_year}}/{{section_code}}                                            நாள்: {{doc_date}}", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.8rem 0"}
            })

            subj = getattr(tpl, "subject_template", None) or "வருவாய் வசூல் சட்டம் 1864 – {{statute_cited}} – {{district_name}} மாவட்டம் - {{taluk_name}} வட்டம் மற்றும் நகரம் – {{defaulter_name}}, {{street_and_locality}} - அரசிற்கு செலுத்த வேண்டிய {{dues_label}} {{total_amount}} - வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்ய கோரியது – உத்திரவிடுதல்."
            ref = getattr(tpl, "reference_template", None) or "{{reference_text}}"

            blocks.append({
                "id": "table_subj_ref",
                "type": "table",
                "rows": [
                    [
                        {"colSpan": 1, "width": "15%", "blocks": [{"id": "p_subj_lbl", "type": "paragraph", "text": "பொருள்:", "runs": [{"text": "பொருள்:", "style": {"fontWeight": "bold"}}], "style": {"textAlign": "left", "fontSize": "11pt"}}]},
                        {"colSpan": 1, "width": "85%", "blocks": [{"id": "p_subj_val", "type": "paragraph", "text": subj, "runs": [{"text": subj}], "style": {"textAlign": "left", "fontSize": "11pt"}}]}
                    ],
                    [
                        {"colSpan": 1, "width": "15%", "blocks": [{"id": "p_ref_lbl", "type": "paragraph", "text": "பார்வை:", "runs": [{"text": "பார்வை:", "style": {"fontWeight": "bold"}}], "style": {"textAlign": "left", "fontSize": "11pt"}}]},
                        {"colSpan": 1, "width": "85%", "blocks": [{"id": "p_ref_val", "type": "paragraph", "text": ref, "runs": [{"text": ref}], "style": {"textAlign": "left", "fontSize": "11pt"}}]}
                    ]
                ]
            })

            blocks.append({
                "id": "p_order_lbl",
                "type": "paragraph",
                "text": "உத்தரவு:",
                "runs": [{"text": "உத்தரவு:", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "left", "fontSize": "11pt", "margin": "0.6rem 0 0.4rem 0"}
            })

            p1 = getattr(tpl, "order_para1_template", None) or "{{district_name}} மாவட்டம், {{taluk_name}} வட்டம் மற்றும் நகரம், {{street_and_locality}} என்ற முகவரியில் {{living_verb}} {{defaulter_name}} என்ற {{entity_label}} அரசிற்கு செலுத்த வேண்டிய {{dues_label}} {{total_amount}}ஐ வருவாய் வசூல் சட்டத்தின் கீழ் வசூலிக்குமாறு பார்வையில் காணும் கடிதத்தில் தெரிவிக்கப்பட்டுள்ளது."
            blocks.append({
                "id": "p_order_para1",
                "type": "paragraph",
                "text": p1,
                "runs": [{"text": p1}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "textIndent": "2rem", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.6rem 0", "lineHeight": "1.6"}
            })

            p2 = getattr(tpl, "order_para2_template", None) or "மேற்படி {{defaulter_name}} {{defaulter_suffix}} தொகை {{total_amount}}ஐ வருவாய் நிலை ஆணை எண் 41 மற்றும் வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூல் செய்ய {{taluk_name}} வருவாய் வட்டாட்சியருக்கு அதிகாரம் வழங்கி உத்திரவிடப்படுகிறது."
            blocks.append({
                "id": "p_order_para2",
                "type": "paragraph",
                "text": p2,
                "runs": [{"text": p2}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "textIndent": "2rem", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.6rem 0", "lineHeight": "1.6"}
            })

            p3 = getattr(tpl, "order_para3_template", None) or "எனவே, மேற்படி முகவரியில் {{living_verb}} {{defaulter_name}} என்பவரின் அசையும் மற்றும் அசையா சொத்துகளிலிருந்து அரசிற்கு செலுத்த வேண்டிய {{dues_label}} {{total_amount}}ஐ ({{amount_in_tamil_words}}) வருவாய் வசூல் சட்டப்படி வசூல் செய்து “{{dd_favour_of}}” என்ற பெயரில் வங்கி வரைவோலையாக (Demand Draft) எடுத்து {{dispatch_address}} என்ற அலுவலகத்திற்கு அசலினை அனுப்பி அதன் விவரத்தினை நகல் வங்கி வரைவோலையுடன் இவ்வலுவலகத்திற்கு அனுப்பி வைக்குமாறு {{taluk_name}} வருவாய் வட்டாட்சியருக்கு தெரிவிக்கப்படுகிறது."
            blocks.append({
                "id": "p_order_para3",
                "type": "paragraph",
                "text": p3,
                "runs": [{"text": p3}],
                "style": {"textAlign": "justify", "fontSize": "11pt", "textIndent": "2rem", "fontFamily": "'TAU-Marutham', serif", "margin": "0 0 0.6rem 0", "lineHeight": "1.6"}
            })

            enclosure = getattr(tpl, "enclosure_text", None) or "கடித நகல்"
            blocks.append({
                "id": "p_enclosure",
                "type": "paragraph",
                "text": f"இணைப்பு: {enclosure}",
                "runs": [{"text": f"இணைப்பு: {enclosure}", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "left", "fontSize": "11pt", "margin": "0.8rem 0 0.4rem 0"}
            })

            signatory = getattr(tpl, "signatory_text", None) or "மாவட்ட ஆட்சித் தலைவர்,\nஈரோடு."
            blocks.append({
                "id": "p_signatory",
                "type": "paragraph",
                "text": signatory,
                "runs": [{"text": signatory, "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "right", "fontSize": "11pt", "margin": "1.2rem 0 0.8rem 0"}
            })

            blocks.append({
                "id": "p_to",
                "type": "paragraph",
                "text": "பெறுநர்:\nவருவாய் வட்டாட்சியர்,\n{{taluk_name}}.",
                "runs": [{"text": "பெறுநர்:\nவருவாய் வட்டாட்சியர்,\n{{taluk_name}}.", "style": {"fontWeight": "bold"}}],
                "style": {"textAlign": "left", "fontSize": "11pt", "margin": "0.6rem 0 0.3rem 0"}
            })

            blocks.append({
                "id": "p_copy1",
                "type": "paragraph",
                "text": "நகல்:\nவருவாய் கோட்டாட்சியர்,\n{{RDO}}.",
                "runs": [{"text": "நகல்:\nவருவாய் கோட்டாட்சியர்,\n{{RDO}}."}],
                "style": {"textAlign": "left", "fontSize": "11pt", "margin": "0.3rem 0"}
            })

            blocks.append({
                "id": "p_copy2",
                "type": "paragraph",
                "text": "நகல்:\n{{dispatch_address}}",
                "runs": [{"text": "நகல்:\n{{dispatch_address}}"}],
                "style": {"textAlign": "left", "fontSize": "11pt", "margin": "0.3rem 0"}
            })

            blocks.append({
                "id": "p_copy3",
                "type": "paragraph",
                "text": "நகல்:\n{{defaulter_name}},\n{{street_and_locality}}, {{village}},\n{{taluk_name}} வட்டம், {{district_name}} மாவட்டம்.",
                "runs": [{"text": "நகல்:\n{{defaulter_name}},\n{{street_and_locality}}, {{village}},\n{{taluk_name}} வட்டம், {{district_name}} மாவட்டம்."}],
                "style": {"textAlign": "left", "fontSize": "11pt", "margin": "0.3rem 0"}
            })

        return {
            "filename": getattr(tpl, "file_name", None) or f"{getattr(tpl, 'template_code', 'template')}.docx",
            "revision": 1,
            "page": {"page_width": 595.3, "page_height": 841.9, "top_margin": 54.0, "right_margin": 54.0, "bottom_margin": 54.0, "left_margin": 54.0},
            "blocks": blocks
        }

    def get_layout(self, filename: str) -> Dict[str, Any]:
        """Locates the file in outputs, templates, or uploads and parses layout."""
        file_path = self.storage.find_file(filename)
        return self.docx_to_layout(file_path, filename)
