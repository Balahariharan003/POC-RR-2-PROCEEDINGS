"""
Template Builder for Tamil Nadu District Collector Proceedings.
Generates pre-formatted docx templates with TAU - Marutham font and Jinja2 tags for docxtpl.
Strictly adheres to official Tamil Nadu Government Proceedings (செயல்முறைகள்) layout standards:
1. Centered and bold Header & Prefix ("முன்னிலை: ...") block with sanitized strings.
2. Borderless 2-column metadata table for Reference File Number (Left) and Date (Right).
3. Bold labeled functional blocks ("பொருள்:", "பார்வை:", "உத்தரவு:") with 12pt vertical paragraph spacing.
"""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from config import PROCEEDINGS_TEMPLATE_PATH, PRIMARY_FONT_TAMIL


def set_font_formatting(run, font_name: str = PRIMARY_FONT_TAMIL, size_pt: float = 12.0, bold: bool = False, italic: bool = False, color_rgb: RGBColor = None):
    """Sets Tamil font name, size, bold, and complex-script formatting for TAU-Marutham/Unicode."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    # Ensure EastAsian / Complex Script (CS) font is also set for Tamil Unicode rendering
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>'
    )
    rPr.append(rFonts)


def _add_header(doc, title: str, prefix: str = "முன்னிலை: {{collector_name}}"):
    """
    Heading Blocks (Title & Prefix):
    - Main title header is perfectly CENTER ALIGNED and BOLD.
    - Prefix line ('முன்னிலை: திரு.ச.கந்தசாமி,இ.ஆ.ப.,') is perfectly CENTER ALIGNED directly underneath.
    """
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(title)
    set_font_formatting(r1, PRIMARY_FONT_TAMIL, 13.0, bold=True)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(14)
    r2 = p2.add_run(prefix)
    set_font_formatting(r2, PRIMARY_FONT_TAMIL, 12.0, bold=True)


def _add_roc_date(doc, roc_tag: str, date_tag: str):
    """
    Tracking Meta-Data Row (File Number & Date):
    Borderless 2-column Word table spanning the full width of the text margins.
    Left: Reference File Number (LEFT ALIGNED)
    Right: Date (RIGHT ALIGNED)
    """
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(3.2)
    
    # Fully remove all table borders
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                '<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
                '<w:insideH w:val="none"/><w:insideV w:val="none"/>'
                '</w:tcBorders>'
            )
            tcPr.append(tcBorders)

    # Left cell: File reference number
    p_left = table.cell(0, 0).paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_left.paragraph_format.space_before = Pt(0)
    p_left.paragraph_format.space_after = Pt(12)
    r_left = p_left.add_run(roc_tag)
    set_font_formatting(r_left, PRIMARY_FONT_TAMIL, 11.5, bold=True)

    # Right cell: Date
    p_right = table.cell(0, 1).paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_before = Pt(0)
    p_right.paragraph_format.space_after = Pt(12)
    r_right = p_right.add_run(date_tag)
    set_font_formatting(r_right, PRIMARY_FONT_TAMIL, 11.5, bold=True)


def _add_labeled_para(doc, label: str, text: str):
    """
    Paragraph Blocks (பொருள், பார்வை):
    - Label styled in BOLD
    - Clean vertical space block (12pt paragraph spacing)
    - Strips leading artifacts
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    clean_text = text.lstrip(" |[")
    r_lbl = p.add_run(label)
    set_font_formatting(r_lbl, PRIMARY_FONT_TAMIL, 11.5, bold=True)
    r_val = p.add_run(clean_text)
    set_font_formatting(r_val, PRIMARY_FONT_TAMIL, 11.5, bold=False)


def _add_divider(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("-------")
    set_font_formatting(r, PRIMARY_FONT_TAMIL, 11.0, bold=True)


def _add_heading(doc, title: str = "உத்தரவு:"):
    """
    Heading Block (உத்தரவு:):
    - Label styled in BOLD
    - Clean vertical space block (at least 12pt paragraph spacing)
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title)
    set_font_formatting(r, PRIMARY_FONT_TAMIL, 12.0, bold=True)


def _add_body_para(doc, text: str):
    """Operative body paragraph with first line indent and clean spacing."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.first_line_indent = Inches(0.4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_font_formatting(r, PRIMARY_FONT_TAMIL, 11.5, bold=False)


def _add_enclosure(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(f"இணைப்பு: {text}")
    set_font_formatting(r, PRIMARY_FONT_TAMIL, 11.5, bold=True)


def _add_signatory(doc, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run(title)
    set_font_formatting(r, PRIMARY_FONT_TAMIL, 11.5, bold=True)


def _add_recipients(doc, list_tuples):
    for lbl, val in list_tuples:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r_lbl = p.add_run(f"{lbl}: ")
        set_font_formatting(r_lbl, PRIMARY_FONT_TAMIL, 11.5, bold=True)
        r_val = p.add_run(val)
        set_font_formatting(r_val, PRIMARY_FONT_TAMIL, 11.5, bold=False)


def _add_office_note_section(doc, roc: str, subject: str, reference: str, submission_paras: list, section_code="ஈ2", district="ஈரோடு"):
    doc.add_page_break()
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("//அலுவலகக் குறிப்பு//")
    set_font_formatting(r_title, PRIMARY_FONT_TAMIL, 12.5, bold=True)

    p_roc = doc.add_paragraph()
    p_roc.paragraph_format.space_after = Pt(6)
    r_roc = p_roc.add_run(roc)
    set_font_formatting(r_roc, PRIMARY_FONT_TAMIL, 11.5, bold=True)

    _add_labeled_para(doc, "பொருள்: ", subject)
    _add_labeled_para(doc, "பார்வை: ", reference)
    _add_divider(doc)
    _add_heading(doc, "பணிந்தனுப்பப்படுகிறது:")

    for para in submission_paras:
        _add_body_para(doc, para)

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(20)
    r_sig = p_sig.add_run(f"ஒப்பம்/–\nபிரிவு எழுத்தர் / கண்காணிப்பாளர்\n{section_code} பிரிவு, மாவட்ட ஆட்சியர் அலுவலகம், {district}.")
    set_font_formatting(r_sig, PRIMARY_FONT_TAMIL, 11.0, bold=True)


def build_proceedings_template(output_path: Path = PROCEEDINGS_TEMPLATE_PATH) -> Path:
    """
    Builds the master District Collector Proceedings Word template matching the standard Tamil Nadu Revenue format:
    Part 1: Proceedings (செயல்முறைகள்) + Signatory + Dispatch List
    Part 2: Office Notes (//அலுவலகக் குறிப்பு//) + Section Submission
    """
    doc = docx.Document()

    # Page Margins: Standard 0.9 inch
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # 1. Header (Centered, Bold)
    _add_header(doc, "{{ collector_heading }}", "முன்னிலை: {{ collector_name }}")

    # 2. Tracking Meta-Data Row (Borderless 2-Column Table)
    _add_roc_date(doc, "{{ roc_number }}", "நாள்: {{ proceedings_date }}")

    # 3. Subject (பொருள்:)
    _add_labeled_para(doc, "பொருள்: ", "{{ subject_text }}")

    # 4. Reference (பார்வை:)
    _add_labeled_para(doc, "பார்வை: ", "{{ reference_text }}")

    # Divider
    _add_divider(doc)

    # 5. Order Heading (உத்தரவு:)
    _add_heading(doc, "உத்தரவு:")

    # 6. Operative Paragraphs
    _add_body_para(doc, "{{ order_para1 }}")
    _add_body_para(doc, "{{ order_para2 }}")
    _add_body_para(doc, "{{ order_para3 }}")

    # 7. Enclosure (இணைப்பு:)
    _add_enclosure(doc, "{{ enclosure_text }}")

    # 8. Signatory (Right aligned)
    _add_signatory(doc, "மாவட்ட ஆட்சித் தலைவர்,\n{{ district_name }}.")

    # 9. Recipients (பெறுநர் / நகல்)
    _add_recipients(doc, [
        ("பெறுநர்", "{{ tahsildar_recipient }}"),
        ("நகல்  ", "{{ rdo_recipient }}")
    ])

    # Office Notes (Part 2)
    _add_office_note_section(
        doc,
        roc="{{ roc_number }}",
        subject="{{ subject_text }}",
        reference="{{ reference_text }}",
        submission_paras=[
            "மேற்படி கோரிக்கையின்படி, உரிய சட்ட விதிகளின் கீழ் வசூலிக்கும் பொருட்டு வட்டாட்சியருக்கு ஆணை பிறப்பித்து செயல்முறைக் குறிப்பாணை தயார் செய்யப்பட்டு மாவட்ட ஆட்சித் தலைவர் அவர்களின் ஒப்புதலுக்குப் பணிந்தனுப்பப்படுகிறது."
        ],
        section_code="{{ section_code }}",
        district="{{ district_name }}"
    )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def build_department_templates(templates_dir: Path = None):
    """
    Builds the four department-specific Word templates:
    1. template_customs.docx
    2. template_tnrera.docx
    3. template_mcop.docx
    4. template_warrant.docx
    and master proceedings_template.docx with exact Tamil Nadu Government Proceedings styling.
    """
    if templates_dir is None:
        templates_dir = PROCEEDINGS_TEMPLATE_PATH.parent
    
    templates_dir.mkdir(parents=True, exist_ok=True)

    # 1. Custom Template A: Customs
    doc_c = docx.Document()
    for section in doc_c.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    _add_header(doc_c, "ஈரோடு மாவட்ட ஆட்சித் தலைவர் மற்றும்\nமாவட்ட நிர்வாக நடுவர் அவர்களின் செயல்முறைகள்", "முன்னிலை: {{collector_name}}")
    _add_roc_date(doc_c, "ந.க. {{file_no}}/{{file_year}}/{{section_code}}", "நாள்:        .05.{{file_year}}.")
    
    _add_labeled_para(doc_c, "பொருள்: ", "வருவாய் வசூல் சட்டம் 1864 – சுங்கச் சட்டம் 1962 பிரிவு 142(1)(c)(i) – ஈரோடு மாவட்டம் – {{taluk_name}} வட்டம் - {{defaulter_name}}, {% if iec_no %}(IEC No: {{iec_no}}){% endif %} {{door_no}}, {{street_and_locality}}, {{taluk_name}} – அரசுக்குச் செலுத்த வேண்டிய நிலுவைத் தொகை வசூல் செய்யக் கோருதல் - உத்திரவிடுதல்.")
    _add_labeled_para(doc_c, "பார்வை: ", "1. {{issuing_authority_name}}, கடித F.NO. {{case_file_no}}, நாள் {{order_date}}.\n2. Order in Original No. {{order_in_original_no}}, நாள் {{letter_date}}.")
    
    _add_divider(doc_c)
    _add_heading(doc_c, "உத்தரவு:")
    _add_body_para(doc_c, "{{district_name}} மாவட்டம், {{taluk_name}} வட்டம், {{street_and_locality}}, {{door_no}}, என்ற முகவரியில் {{living_verb}} {{defaulter_name}} {% if iec_no %}(IEC No: {{iec_no}}){% endif %} {{defaulter_suffix}} சுங்கச் சட்டம் 1962-ன்படி அரசுக்குச் செலுத்த வேண்டிய நிலுவைத் தொகை ரூ.{{total_amount}}/- (அசல் ரூ.{{principal_amount}}/- + அபராதம் ரூ.{{penalty_amount}}/-) மற்றும் வட்டியினை தமிழ்நாடு வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்யுமாறு பார்வையில் காணும் உத்தரவின் வாயிலாக தெரிவிக்கப்பட்டுள்ளது.")
    _add_body_para(doc_c, "மேற்படி {{defaulter_name}} {{defaulter_suffix}} தொகை ரூ.{{total_amount}}/- ஐ வருவாய் நிலை ஆணை எண் 41 மற்றம் தமிழ்நாடு வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூல் செய்ய {{taluk_name}} வருவாய் வட்டாட்சியருக்கு அதிகாரம் வழங்கி இதன் மூலம் உத்திரவிடப்படுகிறது.")
    _add_body_para(doc_c, "எனவே, மேற்படி முகவரியில் {{living_verb}} {{defaulter_name}} என்பாரின் {{asset_clause}} மொத்தம் தொகை ரூ.{{total_amount}}/- ({{amount_in_tamil_words}}) மற்றும் உரிய வட்டியினை வசூல் செய்து “{{dd_favour_of}}“ என்ற பெயரில் வங்கி வரைவோலையாக (Demand Draft - Head of Account: {{head_of_account}}) எடுத்து {{dispatch_address}} என்ற அலுவலகத்திற்கு அசலினை அனுப்பி அதன் விவரத்தினை நகல் வங்கி வரைவோலையுடன் இவ்வலுவலகத்திற்கு அனுப்பி வைக்குமாறு {{taluk_name}} வருவாய் வட்டாட்சியருக்கு தெரிவிக்கப்படுகிறது.")
    
    _add_enclosure(doc_c, "கடித நகல்")
    _add_signatory(doc_c, "மாவட்ட ஆட்சித் தலைவர்,\n{{district_name}}.")
    _add_recipients(doc_c, [
        ("பெறுநர்", "வருவாய் வட்டாட்சியர், {{taluk_name}}."),
        ("நகல்  ", "வருவாய் கோட்டாட்சியர், {{district_name}}."),
        ("நகல்  ", "{{dispatch_address}}"),
        ("நகல்  ", "{{defaulter_name}}, {{door_no}}, {{street_and_locality}}, {{taluk_name}} - {{pincode}}.")
    ])
    
    _add_office_note_section(
        doc_c,
        roc="ந.க. {{file_no}}/{{file_year}}/{{section_code}}",
        subject="வருவாய் வசூல் சட்டம் 1864 – சுங்கச் சட்டம் 1962 பிரிவு 142(1)(c)(i) – சென்னை சுங்கத்துறை நிலுவைத் தொகை வசூலித்தல் – ஆணை பிறப்பித்தல் – சார்பு.",
        reference="சென்னை சுங்கத்துறை ஆணையரக கடிதம் F.NO. {{case_file_no}}, நாள் {{order_date}}.",
        submission_paras=[
            "பார்வையில் கண்டுள்ள கடிதத்தில், {{defaulter_name}} நிறுவனம் செலுத்த வேண்டிய சுங்கத் தீர்வை மற்றும் அபராதத் தொகை ரூ.{{total_amount}}/-யினை வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூலிக்கக் கோரப்பட்டுள்ளது.",
            "இதன்மீது நடவடிக்கை மேற்கொள்ளும் வகையில், {{taluk_name}} வட்டாட்சியருக்கு தமிழ்நாடு வருவாய் வசூல் சட்டம் 1864 பிரிவு 5 மற்றும் சுங்கச் சட்டம் 1962 பிரிவு 142(1)(c)(i)-ன் கீழ் ஆணை பிறப்பித்து செயல்முறைக் குறிப்பாணை தயார் செய்யப்பட்டு மாவட்ட ஆட்சித் தலைவர் அவர்களின் ஒப்புதலுக்குப் பணிந்தனுப்பப்படுகிறது."
        ],
        section_code="{{section_code}}",
        district="{{district_name}}"
    )
    doc_c.save(str(templates_dir / "template_customs.docx"))

    # 2. Custom Template B: TNRERA
    doc_r = docx.Document()
    for section in doc_r.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    _add_header(doc_r, "ஈரோடு மாவட்ட ஆட்சித் தலைவர் மற்றும்\nமாவட்ட நிர்வாக நடுவர் அவர்களின் செயல்முறைகள்", "முன்னிலை: {{collector_name}}")
    _add_roc_date(doc_r, "ந.க. {{file_no}}/{{file_year}}/{{section_code}}", "நாள்:        .05.{{file_year}}.")
    _add_labeled_para(doc_r, "பொருள்: ", "வருவாய் வசூல் சட்டம் 1864 – தமிழ்நாடு ரியல் எஸ்டேட் (முறைப்படுத்துதல் மற்றும் மேம்படுத்துதல்) சட்டம் 2016 பிரிவு 40(1) – {{district_name}} மாவட்டம் – {{taluk_name}} வட்டம் - {{defaulter_details[0].name}}, {{defaulter_details[0].door_no}}, {{defaulter_details[0].street_and_locality}} – அபராதத் தொகை வசூல் செய்யக் கோருதல் - உத்திரவிடுதல்.")
    _add_labeled_para(doc_r, "பார்வை: ", "{{reference_details.issuing_authority_name}}, கடித எண். {{reference_details.case_or_file_no}}, நாள் {{reference_details.order_date}}.")
    _add_divider(doc_r)
    _add_heading(doc_r, "உத்தரவு:")
    _add_body_para(doc_r, "{{district_name}} மாவட்டம், {{taluk_name}} வட்டம், {{defaulter_details[0].street_and_locality}}, {{defaulter_details[0].door_no}}, என்ற முகவரியில் {{living_verb}} {{defaulter_details[0].name}} {{defaulter_suffix}} தமிழ்நாடு ரியல் எஸ்டேட் சட்டம் 2016 பிரிவு 40(1)-ன் படி விதிக்கப்பட்ட அபராதத் தொகை ரூ.{{financials.total_amount}}/- ஐ தமிழ்நாடு வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்யுமாறு பார்வையில் காணும் உத்தரவின் வாயிலாக தெரிவிக்கப்பட்டுள்ளது.")
    _add_body_para(doc_r, "மேற்படி {{defaulter_details[0].name}} {{defaulter_suffix}} தொகை ரூ.{{financials.total_amount}}/- ஐ வருவாய் நிலை ஆணை எண் 41 மற்றம் வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூல் செய்ய {{taluk_name}} வருவாய் வட்டாட்சியருக்கு அதிகாரம் வழங்கி இதன் மூலம் உத்திரவிடப்படுகிறது.")
    _add_body_para(doc_r, "எனவே, மேற்படி முகவரியில் {{living_verb}} {{defaulter_details[0].name}} என்பாரின் {{asset_clause}} மொத்தம் தொகை ரூ.{{financials.total_amount}}/- ({{financials.amount_in_tamil_words}}) வசூல் செய்து “{{payment_instructions.dd_favour_of}}“ என்ற பெயரில் வங்கி வரைவோலையாக (Demand Draft) எடுத்து {{payment_instructions.dispatch_address}} என்ற அலுவலகத்திற்கு அசலினை அனுப்பி அதன் விவரத்தினை நகல் வங்கி வரைவோலையுடன் இவ்வலுவலகத்திற்கு அனுப்பி வைக்குமாறு {{taluk_name}} வருவாய் வட்டாட்சியருக்கு தெரிவிக்கப்படுகிறது.")
    _add_enclosure(doc_r, "கடித நகல்")
    _add_signatory(doc_r, "மாவட்ட ஆட்சித் தலைவர்,\n{{district_name}}.")
    _add_recipients(doc_r, [
        ("பெறுநர்", "வருவாய் வட்டாட்சியர், {{taluk_name}}."),
        ("நகல்  ", "வருவாய் கோட்டாட்சியர், {{district_name}}."),
        ("நகல்  ", "{{payment_instructions.dispatch_address}}"),
        ("நகல்  ", "{{defaulter_details[0].name}}, {{defaulter_details[0].door_no}}, {{defaulter_details[0].street_and_locality}}, {{taluk_name}} - {{defaulter_details[0].pincode}}.")
    ])
    _add_office_note_section(
        doc_r,
        roc="ந.க. {{file_no}}/{{file_year}}/{{section_code}}",
        subject="வருவாய் வசூல் சட்டம் 1864 – தமிழ்நாடு ரியல் எஸ்டேட் சட்டம் 2016 பிரிவு 40(1) – TNRERA நிலுவைத் தொகை வசூலித்தல் – ஆணை பிறப்பித்தல் – சார்பு.",
        reference="TNRERA ஆணையரக கடிதம் {{reference_details.case_or_file_no}}, நாள் {{reference_details.order_date}}.",
        submission_paras=[
            "பார்வையில் கண்டுள்ள கடிதத்தில், {{defaulter_details[0].name}} செலுத்த வேண்டிய அபராதத் தொகை ரூ.{{financials.total_amount}}/-யினை வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூலிக்கக் கோரப்பட்டுள்ளது.",
            "இதன்மீது நடவடிக்கை மேற்கொள்ளும் வகையில், {{taluk_name}} வட்டாட்சியருக்கு தமிழ்நாடு வருவாய் வசூல் சட்டம் 1864 பிரிவு 5 மற்றும் TNRERA சட்டம் 2016 பிரிவு 40(1)-ன் கீழ் ஆணை பிறப்பித்து செயல்முறைக் குறிப்பாணை தயார் செய்யப்பட்டு மாவட்ட ஆட்சித் தலைவர் அவர்களின் ஒப்புதலுக்குப் பணிந்தனுப்பப்படுகிறது."
        ],
        section_code="{{section_code}}",
        district="{{district_name}}"
    )
    doc_r.save(str(templates_dir / "template_tnrera.docx"))

    # 3. Custom Template C: MCOP
    doc_m = docx.Document()
    for section in doc_m.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    _add_header(doc_m, "ஈரோடு மாவட்ட ஆட்சித் தலைவர் மற்றும்\nமாவட்ட நிர்வாக நடுவர் அவர்களின் செயல்முறைகள்", "முன்னிலை: {{collector_name}}")
    _add_roc_date(doc_m, "ந.க. {{file_no}}/{{file_year}}/{{section_code}}", "நாள்:        .05.{{file_year}}.")
    _add_labeled_para(doc_m, "பொருள்: ", "வருவாய் வசூல் சட்டம் 1864 – மோட்டார் வாகனச் சட்டம் 1988 – {{district_name}} மாவட்டம் – {{taluk_name}} வட்டம் - {{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}}, {% endif %}{{defaulter_details[0].door_no}}, {{defaulter_details[0].street_and_locality}}, {{taluk_name}} - மோட்டார் வாகனச் சட்டம் 1988 பிரிவு 174 - {{reference_details.issuing_authority_name}} - {{reference_details.ia_or_mp_no}} -ன் {{reference_details.case_or_file_no}} -இன் படி தொகை ரூ.{{financials.total_amount}}/- {% if financials.interest_rate %}ஐ {{financials.interest_rate}}% வட்டியுடன் {% endif %}வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்ய கோரியது - உத்திரவிடுதல்.")
    _add_labeled_para(doc_m, "பார்வை: ", "{{reference_details.issuing_authority_name}}, {{reference_details.case_or_file_no}}, உத்தரவு, நாள் {{reference_details.order_date}}.")
    _add_divider(doc_m)
    _add_heading(doc_m, "உத்தரவு:")
    _add_body_para(doc_m, "{{district_name}} மாவட்டம், {{taluk_name}} வட்டம், {{defaulter_details[0].street_and_locality}}, {{defaulter_details[0].door_no}}, என்ற முகவரியில் {{living_verb}} {{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}} {% endif %}{{defaulter_suffix}} மோட்டார் வாகனச் சட்டம் 1988 பிரிவு 174, {{reference_details.issuing_authority_name}} {{reference_details.ia_or_mp_no}} -ன் {{reference_details.case_or_file_no}} -இன் படி தொகை ரூ.{{financials.total_amount}}/- {% if financials.interest_rate %}ஐ {{financials.interest_rate}}% வட்டியுடன் {% endif %}வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்யுமாறு பார்வையில் காணும் உத்தரவின் வாயிலாக தெரிவிக்கப்பட்டுள்ளது.")
    _add_body_para(doc_m, "மேற்படி {{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}} {% endif %}{{defaulter_suffix}} தொகை ரூ.{{financials.total_amount}}/- {% if financials.interest_rate %}ஐ {{financials.interest_rate}}% வட்டியுடன் {% endif %}வருவாய் நிலை ஆணை எண் 41 மற்றம் வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூல் செய்ய {{taluk_name}} வருவாய் வட்டாட்சியருக்கு அதிகாரம் வழங்கி இதன் மூலம் உத்திரவிடப்படுகிறது.")
    _add_body_para(doc_m, "எனவே, மேற்படி முகவரியில் {{living_verb}} {{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}} {% endif %}என்பவரின் {{asset_clause}} {{split_clause}}மொத்தம் தொகை ரூ.{{financials.total_amount}}/- ({{financials.amount_in_tamil_words}}) {% if financials.interest_start_date %}மற்றும் கடந்த {{financials.interest_start_date}} முதல் தொகை செலுத்தும் நாள் வரையில் {{financials.interest_rate}}% வட்டித்தொகையுடன் சேர்த்து {% endif %}வருவாய் வசூல் சட்டப்படி வசூல் செய்து “{{payment_instructions.dd_favour_of}}“ என்ற பெயரில் வங்கி வரைவோலையாக (Demand Draft) எடுத்து {{payment_instructions.dispatch_address}} என்ற அலுவலகத்திற்கு அசலினை அனுப்பி அதன் விவரத்தினை நகல் வங்கி வரைவோலையுடன் {{reference_details.issuing_authority_name}} என்ற நீதிமன்றத்திற்கும் மற்றும் இவ்வலுவலகத்திற்கும் அனுப்பி வைக்குமாறு {{taluk_name}} வருவாய் வட்டாட்சியருக்கு தெரிவிக்கப்படுகிறது.")
    _add_enclosure(doc_m, "கடித நகல்")
    _add_signatory(doc_m, "மாவட்ட ஆட்சித் தலைவர்,\n{{district_name}}.")
    _add_recipients(doc_m, [
        ("பெறுநர்", "வருவாய் வட்டாட்சியர், {{taluk_name}}."),
        ("நகல்  ", "வருவாய் கோட்டாட்சியர், {{district_name}}."),
        ("நகல்  ", "{{payment_instructions.dispatch_address}}"),
        ("நகல்  ", "{{reference_details.issuing_authority_name}}"),
        ("நகல்  ", "{{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}}, {% endif %}{{defaulter_details[0].door_no}}, {{defaulter_details[0].street_and_locality}}, {{taluk_name}} - {{defaulter_details[0].pincode}}.")
    ])
    _add_office_note_section(
        doc_m,
        roc="ந.க. {{file_no}}/{{file_year}}/{{section_code}}",
        subject="வருவாய் வசூல் சட்டம் 1864 – மோட்டார் வாகனச் சட்டம் 1988 – {{district_name}} மாவட்டம் – {{taluk_name}} வட்டம் - {{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}}, {% endif %}{{defaulter_details[0].door_no}}, {{defaulter_details[0].street_and_locality}}, {{taluk_name}} - மோட்டார் வாகனச் சட்டம் 1988 பிரிவு 174 - {{reference_details.issuing_authority_name}} - {{reference_details.ia_or_mp_no}} -ன் {{reference_details.case_or_file_no}} -இன் படி தொகை ரூ.{{financials.total_amount}}/- {% if financials.interest_rate %}ஐ {{financials.interest_rate}}% வட்டியுடன் {% endif %}வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்ய கோரியது - உத்திரவிடுதல்.",
        reference="{{reference_details.issuing_authority_name}}, {{reference_details.case_or_file_no}}, உத்தரவு, நாள் {{reference_details.order_date}}.",
        submission_paras=[
            "{{district_name}} மாவட்டம், {{taluk_name}} வட்டம், {{defaulter_details[0].street_and_locality}}, {{defaulter_details[0].door_no}}, என்ற முகவரியில் {{living_verb}} {{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}}, {% endif %}{{defaulter_suffix}} மோட்டார் வாகனச் சட்டம் 1988 பிரிவு 174, {{reference_details.issuing_authority_name}} {{reference_details.ia_or_mp_no}} -ன் {{reference_details.case_or_file_no}} -இன் படி தொகை ரூ.{{financials.total_amount}}/- {% if financials.interest_rate %}ஐ {{financials.interest_rate}}% வட்டியுடன் {% endif %}வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்யுமாறு பார்வை 1இல் காணும் உத்தரவின் வாயிலாக தெரிவிக்கப்பட்டுள்ளது.",
            "மேற்படி முகவரியில் {{living_verb}} {{defaulter_details[0].name}}, {% if defaulter_details[0].father_or_spouse_name %}{{defaulter_details[0].father_or_spouse_name}} {% endif %}என்பவரின் {{asset_clause}} {{split_clause}}மொத்தம் தொகை ரூ.{{financials.total_amount}}/- ({{financials.amount_in_tamil_words}}) {% if financials.interest_start_date %}மற்றும் கடந்த {{financials.interest_start_date}} முதல் தொகை செலுத்தும் நாள் வரையில் {{financials.interest_rate}}% வட்டித்தொகையுடன் சேர்த்து {% endif %}வருவாய் வசூல் சட்டப்படி வசூல் செய்து “{{payment_instructions.dd_favour_of}}“ என்ற பெயரில் வங்கி வரைவோலையாக (Demand Draft) எடுத்து {{payment_instructions.dispatch_address}} என்ற அலுவலகத்திற்கு அசலினை அனுப்பி அதன் விவரத்தினை நகல் வங்கி வரைவோலையுடன் {{reference_details.issuing_authority_name}} என்ற நீதிமன்றத்திற்கும் மற்றும் இவ்வலுவலகத்திற்கும் அனுப்பி வைக்குமாறு {{taluk_name}} வருவாய் வட்டாட்சியருக்கு தெரிவிக்கலாம்.",
            "எனவே மேற்படி தொகையை வருவாய் வசூல் சட்டம் 1864 பிரிவு 5 / வருவாய் நிலை ஆணை எண்.41-ன் கீழ் வசூல் செய்ய {{taluk_name}} வருவாய் வட்டாட்சியருக்கு அதிகாரம் வழங்கி இதன் மூலம் உத்திரவிடலாம்.",
            "உத்திரவினை எதிர்நோக்கி செயல்முறை வரைவு ஒப்புதலுக்காக மாவட்ட ஆட்சித்தலைவர் அவர்களுக்கு பணிவுடன் சமர்ப்பிக்கப்படுகிறது."
        ],
        section_code="{{section_code}}",
        district="{{district_name}}"
    )
    doc_m.save(str(templates_dir / "template_mcop.docx"))

    # 4. Custom Template D: Warrant
    doc_w = docx.Document()
    for section in doc_w.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    _add_header(doc_w, "ஈரோடு மாவட்ட ஆட்சித் தலைவர் மற்றும்\nமாவட்ட நிர்வாக நடுவர் அவர்களின் செயல்முறைகள்", "முன்னிலை: {{collector_name}}")
    _add_roc_date(doc_w, "ந.க. {{file_no}}/{{file_year}}/{{section_code}}", "நாள்:        .05.{{file_year}}.")
    _add_labeled_para(doc_w, "பொருள்: ", "வருவாய் வசூல் சட்டம் 1864 – குற்றவியல் நடைமுறைச் சட்டம் – வாரண்ட் வசூலித்தல் – {{district_name}} மாவட்டம் – {{taluk_name}} வட்டம் - {{defaulter_details[0].name}}, {{defaulter_details[0].door_no}}, {{defaulter_details[0].street_and_locality}}, {{taluk_name}} – நிலுவைத் தொகை ரூ.{{financials.total_amount}}/- வசூல் செய்யக் கோருதல் - உத்திரவிடுதல்.")
    _add_labeled_para(doc_w, "பார்வை: ", "{{reference_details.issuing_authority_name}}, வாரண்ட் ஆணை {{reference_details.case_or_file_no}}, நாள் {{reference_details.order_date}}.")
    _add_divider(doc_w)
    _add_heading(doc_w, "உத்தரவு:")
    _add_body_para(doc_w, "{{district_name}} மாவட்டம், {{taluk_name}} வட்டம், {{defaulter_details[0].street_and_locality}}, {{defaulter_details[0].door_no}}, என்ற முகவரியில் {{living_verb}} {{defaulter_details[0].name}} {{defaulter_suffix}} நீதிமன்ற வாரண்ட் ஆணைப்படி தொகை ரூ.{{financials.total_amount}}/- ஐ தமிழ்நாடு வருவாய் வசூல் சட்டத்தின் கீழ் வசூல் செய்யுமாறு பார்வையில் காணும் உத்தரவின் வாயிலாக தெரிவிக்கப்பட்டுள்ளது.")
    _add_body_para(doc_w, "மேற்படி {{defaulter_details[0].name}} {{defaulter_suffix}} தொகை ரூ.{{financials.total_amount}}/- ஐ வருவாய் நிலை ஆணை எண் 41 மற்றம் வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூல் செய்ய {{taluk_name}} வருவாய் வட்டாட்சியருக்கு அதிகாரம் வழங்கி இதன் மூலம் உத்திரவிடப்படுகிறது.")
    _add_body_para(doc_w, "எனவே, மேற்படி முகவரியில் {{living_verb}} {{defaulter_details[0].name}} என்பாரின் {{asset_clause}} {{split_clause}}மொத்தம் தொகை ரூ.{{financials.total_amount}}/- ({{financials.amount_in_tamil_words}}) வசூல் செய்து “{{payment_instructions.dd_favour_of}}“ என்ற பெயரில் வங்கி வரைவோலையாக (Demand Draft) எடுத்து {{payment_instructions.dispatch_address}} என்ற அலுவலகத்திற்கு அசலினை அனுப்பி அதன் விவரத்தினை நகல் வங்கி வரைவோலையுடன் இவ்வலுவலகத்திற்கு அனுப்பி வைக்குமாறு {{taluk_name}} வருவாய் வட்டாட்சியருக்கு தெரிவிக்கப்படுகிறது.")
    _add_enclosure(doc_w, "வாரண்ட் நகல்")
    _add_signatory(doc_w, "மாவட்ட ஆட்சித் தலைவர்,\n{{district_name}}.")
    _add_recipients(doc_w, [
        ("பெறுநர்", "வருவாய் வட்டாட்சியர், {{taluk_name}}."),
        ("நகல்  ", "வருவாய் கோட்டாட்சியர், {{district_name}}."),
        ("நகல்  ", "{{payment_instructions.dispatch_address}}"),
        ("நகல்  ", "{{defaulter_details[0].name}}, {{defaulter_details[0].door_no}}, {{defaulter_details[0].street_and_locality}}, {{taluk_name}} - {{defaulter_details[0].pincode}}.")
    ])
    _add_office_note_section(
        doc_w,
        roc="ந.க. {{file_no}}/{{file_year}}/{{section_code}}",
        subject="வருவாய் வசூல் சட்டம் 1864 – குற்றவியல் நடைமுறைச் சட்டம் – வாரண்ட் தொகையினை வசூலித்தல் – ஆணை பிறப்பித்தல் – சார்பு.",
        reference="{{reference_details.issuing_authority_name}}, வாரண்ட் ஆணை {{reference_details.case_or_file_no}}, நாள் {{reference_details.order_date}}.",
        submission_paras=[
            "பார்வையில் கண்டுள்ள கடிதத்தில், {{defaulter_details[0].name}} செலுத்த வேண்டிய வாரண்ட் தொகை ரூ.{{financials.total_amount}}/-யினை வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் வசூலிக்கக் கோரப்பட்டுள்ளது.",
            "இதன்மீது நடவடிக்கை மேற்கொள்ளும் வகையில், {{taluk_name}} வட்டாட்சியருக்கு தமிழ்நாடு வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் ஆணை பிறப்பித்து செயல்முறைக் குறிப்பாணை தயார் செய்யப்பட்டு மாவட்ட ஆட்சித் தலைவர் அவர்களின் ஒப்புதலுக்குப் பணிந்தனுப்பப்படுகிறது."
        ],
        section_code="{{section_code}}",
        district="{{district_name}}"
    )
    doc_w.save(str(templates_dir / "template_warrant.docx"))

    # Also build master proceedings_template.docx
    build_proceedings_template(templates_dir / "proceedings_template.docx")
    print("All 4 department templates + master proceedings template built successfully.")


if __name__ == "__main__":
    build_department_templates()
