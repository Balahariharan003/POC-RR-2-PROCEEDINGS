"""
Unit Test Suite for Phase P3: Document Synthesis & Typography
=============================================================
Verifies template rendering, slot filling, and TAU-Marutham font styling.
"""

import pytest
import asyncio
from pathlib import Path
import docx

from app.services.document_service import (
    process_office_note,
    process_proceedings,
    process_memorandum,
    process_warrant,
    render_text_stream_to_docx
)
from app.core.config import settings


@pytest.mark.asyncio
async def test_process_proceedings_rendering():
    case = {
        "case_file_no": "1248/2026",
        "department_type": "CUSTOMS",
        "statute_cited": "Customs Act 1962 Sec 142(1)(c)(ii)",
        "taluk_name": "பெருந்துறை",
        "district_name": "ஈரோடு",
        "defaulter_name": "M/s Sri Venkateshwara Tex",
        "door_no": "Plot No 45",
        "street_and_locality": "SIPCOT Industrial Complex",
        "village": "Perundurai",
        "pincode": "638052",
        "principal_amount": 250000.0,
        "penalty_amount": 50000.0,
        "interest_amount": 15000.0,
        "total_recoverable_amount": 315000.0,
        "review_flags": [],
        "synthesized_paragraphs": {
            "order_para1": "எதிர்மனுதாரர் M/s Sri Venkateshwara Tex நிறுவனம் ரூ. 3,15,000/- நிலுவை வைத்துள்ளார்.",
            "order_para2": "வருவாய் வசூல் சட்டம் 1864 பிரிவு 5-ன் கீழ் இந்நிலுவை தொகையை வசூலிக்க உத்தரவிடப்படுகிறது.",
            "order_para3": "வசூலிக்கப்பட்ட தொகையை உரிய கணக்கில் வரவு வைக்க வேண்டும்."
        }
    }
    
    meta = {
        "nk_no": "1248/2026",
        "year": "2026",
        "doc_date": ".03.2026"
    }
    
    docx_path_str = await process_proceedings("", 9901, case=case, meta=meta)
    doc_path = Path(docx_path_str)
    assert doc_path.exists()
    
    doc = docx.Document(str(doc_path))
    full_text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    
    # Assert no unfilled slot markers remain
    assert "«" not in full_text
    assert "»" not in full_text
    assert "<<" not in full_text
    assert ">>" not in full_text
    
    # Assert defaulter and amount are rendered
    assert "M/s Sri Venkateshwara Tex" in full_text
    assert "3,15,000" in full_text


@pytest.mark.asyncio
async def test_warrant_only_for_maintenance():
    """Warrant worker must return None for non-maintenance cases."""
    customs_case = {
        "case_file_no": "1248/2026",
        "department_type": "CUSTOMS",
        "taluk_name": "பெருந்துறை",
        "defaulter_name": "ABC Corp",
        "total_recoverable_amount": 100000.0
    }
    warrant_res = await process_warrant("", 9902, case=customs_case)
    assert warrant_res is None


@pytest.mark.asyncio
async def test_tau_marutham_font_enforcement(tmp_path):
    """Verifies that generated runs use TAU-Marutham font."""
    test_docx = tmp_path / "test_font.docx"
    content = "ஈரோடு மாவட்ட ஆட்சித் தலைவர் செயல்முறைகள்\n\nந.க. 1248/2026/ஈ2\n\nஉத்தரவு பிறப்பிக்கப்படுகிறது."
    await render_text_stream_to_docx(content, test_docx, doc_type="PROCEEDINGS")
    
    assert test_docx.exists()
    doc = docx.Document(str(test_docx))
    for p in doc.paragraphs:
        for r in p.runs:
            assert r.font.name == settings.PRIMARY_FONT_TAMIL
