"""
Unit and Integration Tests for Editor Service and Document Layout Parsing.
Verifies:
1. Block ID sequential alignment between parser and exporter.
2. Table cell paragraph and top-level paragraph edits.
3. Word document layout extraction and DB template model conversion.
"""

import pytest
import docx
from pathlib import Path
import tempfile

from app.services.editor_service import EditorService
from app.services.document_service import DocumentService
from app.domain.schemas.legal_entities import (
    ExtractedLegalEntities,
    DefaulterDetail,
    FinancialDetails,
    ReferenceDetails,
    PaymentInstructions,
)


@pytest.mark.asyncio
async def test_editor_layout_and_export_synchronization():
    from app.services.document_service import process_proceedings
    editor_service = EditorService()

    case = {
        "case_file_no": "F.NO. 516/2024-ARC",
        "department_type": "CUSTOMS",
        "statute_cited": "சுங்கச் சட்டம் 1962 பிரிவு 142(1)(c)(ii)",
        "taluk_name": "கொடுமுடி",
        "district_name": "ஈரோடு",
        "defaulter_name": "திரு.T.P.ராமலிங்கம்",
        "door_no": "14",
        "street_and_locality": "பவானி மெயின் ரோடு",
        "village": "கொடுமுடி",
        "pincode": "638151",
        "principal_amount": 173308.0,
        "penalty_amount": 9000.0,
        "total_recoverable_amount": 182308.0,
        "review_flags": [],
        "synthesized_paragraphs": {
            "order_para1": "சுங்கவரி நிலுவைத் தொகை ரூ. 1,82,308/- (சுங்கவரி ரூ.1,73,308/- + அபராதத் தொகை ரூ.9,000/-) வசூலிக்க கோரப்பட்டுள்ளது.",
            "order_para2": "வருவாய் நிலை ஆணை எண் 41 மற்றும் பிரிவு 5-ன் கீழ் அதிகாரம் வழங்கப்படுகிறது.",
            "order_para3": "வங்கி வரைவோலையாக எடுத்து உரிய அலுவலகத்திற்கு அனுப்பி வைக்க வேண்டும்."
        }
    }
    meta = {
        "nk_no": "516/2024",
        "year": "2026",
        "doc_date": "22.01.2026"
    }

    # 1. Generate real DOCX
    docx_path_str = await process_proceedings("", 9902, case=case, meta=meta)
    docx_path = Path(docx_path_str)
    assert docx_path.exists()


    # 2. Parse into layout
    layout = editor_service.docx_to_layout(Path(docx_path), Path(docx_path).name)
    blocks = layout.get("blocks", [])
    assert len(blocks) > 0

    # Verify paragraph and table structure
    para_blocks = [b for b in blocks if b.get("type") == "paragraph"]
    tbl_blocks = [b for b in blocks if b.get("type") == "table"]
    assert len(para_blocks) > 0
    assert len(tbl_blocks) > 0

    # 3. Simulate user edits on a top-level paragraph AND a table cell paragraph
    p1 = para_blocks[0]
    tbl = tbl_blocks[0]
    cell_p = tbl["rows"][0][1]["blocks"][0]

    edits = {
        p1["id"]: "ஈரோடு மாவட்ட ஆட்சித் தலைவர் அவர்களின் திருத்தப்பட்ட செயல்முறைகள்",
        cell_p["id"]: "வருவாய் வசூல் சட்டம் – திருத்தப்பட்ட பொருள் விவரம்",
    }

    # 4. Apply edits into a new docx and verify both edits are present
    doc = docx.Document(str(docx_path))
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    p_count = 0
    t_count = 0
    for elem in doc.element.body:
        if elem.tag.endswith("p"):
            p = Paragraph(elem, doc)
            p_count += 1
            pid = f"p_{p_count}"
            if pid in edits and edits[pid] is not None:
                p.text = edits[pid]
        elif elem.tag.endswith("tbl"):
            t_count += 1
            table = Table(elem, doc)
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    for cp_idx, cp in enumerate(cell.paragraphs):
                        p_count += 1
                        c_pid = f"p_t{t_count}_r{r_idx}_c{c_idx}_{cp_idx+1}"
                        if c_pid in edits and edits[c_pid] is not None:
                            cp.text = edits[c_pid]

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp.name)
        saved_doc = docx.Document(tmp.name)
        all_text = " ".join([p.text for p in saved_doc.paragraphs] + [c.text for t in saved_doc.tables for r in t.rows for c in r.cells])
        assert "திருத்தப்பட்ட செயல்முறைகள்" in all_text
        assert "திருத்தப்பட்ட பொருள் விவரம்" in all_text


def test_office_note_layout_generation():
    editor_service = EditorService()

    class MockTemplate:
        template_code = "office_note_default"
        category = "INTERNAL_NOTE"
        heading_prefix = "// அலுவலகக் குறிப்பு //"
        subject_template = "பொருள் மாதிரி"
        reference_template = "பார்வை மாதிரி"
        order_para1_template = "பத்தி 1"
        order_para2_template = "பத்தி 2"
        order_para3_template = "பத்தி 3"
        signatory_text = "ஒப்புதலுக்காக சமர்ப்பிக்கப்படுகிறது."
        file_name = "office_note_default.docx"
        template_data = None

    layout = editor_service.template_model_to_layout(MockTemplate())
    assert layout["filename"] == "office_note_default.docx"
    assert len(layout["blocks"]) >= 5
    heading_block = [b for b in layout["blocks"] if b["id"] == "p_heading"][0]
    assert "// அலுவலகக் குறிப்பு //" in heading_block["text"]
